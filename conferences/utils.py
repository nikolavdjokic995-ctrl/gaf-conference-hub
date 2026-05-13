from docx import Document
from lxml import etree

import zipfile
import tempfile
import shutil
import os
import re


AUTHOR_KEYWORDS = (
    "affiliation", "email", "e-mail", "university", "faculty", "department",
    "institute", "college", "school", "laboratory", "centre", "center",
    "author", "coauthor", "co-author", "corresponding author", "orcid",
)
STOP_SECTION_WORDS = ("abstract", "keywords", "key words", "introduction")


def _iter_paragraphs(parent):
    """Yield normal paragraphs and paragraphs nested inside tables."""
    for para in parent.paragraphs:
        yield para

    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _blank_paragraph(para):
    """Clear paragraph text while keeping the document structure valid."""
    para.text = ""


def _looks_like_identity_line(text, before_main_text=False):
    clean = " ".join(text.split())
    lower = clean.lower()

    if not clean:
        return False

    if any(word in lower for word in AUTHOR_KEYWORDS):
        return True

    if "@" in clean or "http" in lower or "www." in lower:
        return True

    # Common affiliation markers: superscript numbers are often flattened by python-docx.
    if "," in clean and any(char.isdigit() for char in clean) and len(clean) < 300:
        return True

    # Author/co-author line near the top: short comma-separated list of names.
    if before_main_text and "," in clean and len(clean) < 180:
        words = [w for w in re.split(r"[\s,;]+", clean) if w]
        capitalized = sum(1 for w in words if w[:1].isupper())
        if len(words) >= 2 and capitalized >= max(2, len(words) // 2):
            return True

    return False


def _strip_docx_xml_text(docx_path):
    """Remove identity text from DOCX XML parts not fully exposed by python-docx."""
    tmp_fd, tmp_path = tempfile.mkstemp(suffix=".docx")
    os.close(tmp_fd)

    text_pattern = re.compile(r"(<w:t(?:\s+[^>]*)?>)(.*?)(</w:t>)", re.DOTALL)

    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)

            if item.filename in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}:
                xml = data.decode("utf-8", errors="ignore")
                xml = text_pattern.sub(r"\1\3", xml)
                data = xml.encode("utf-8")

            zout.writestr(item, data)

    shutil.move(tmp_path, docx_path)


def anonymize_docx(source_path, target_path):
    doc = Document(source_path)

    props = doc.core_properties
    for field in ["author", "last_modified_by", "comments", "title", "subject", "category", "keywords"]:
        try:
            setattr(props, field, "")
        except Exception:
            pass

    paragraphs = list(doc.paragraphs)

    abstract_index = None
    for i, para in enumerate(paragraphs):
        if para.text.strip().lower().startswith("abstract"):
            abstract_index = i
            break

    if abstract_index is not None:
        # Briše samo nekoliko pasusa neposredno iznad Abstract-a: autori/coauthors
        removed = 0
        for j in range(abstract_index - 1, -1, -1):
            text = paragraphs[j].text.strip()

            if not text:
                continue

            if "title of the paper" in text.lower():
                break

            if removed >= 4:
                break

            paragraphs[j].text = ""
            removed += 1

    for section in doc.sections:
        for para in section.header.paragraphs:
            para.text = ""
        for para in section.footer.paragraphs:
            para.text = ""

    temp_docx = target_path + ".tmp.docx"
    doc.save(temp_docx)

    # Briše tekst fusnota, ali ne dira separator/crtu fusnote
    with zipfile.ZipFile(temp_docx, "r") as zin:
        with zipfile.ZipFile(target_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)

                if item.filename == "word/footnotes.xml":
                    root = etree.fromstring(data)
                    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

                    for footnote in root.findall("w:footnote", ns):
                        fid = footnote.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id")
                        if fid not in ("-1", "0"):
                            for t in footnote.findall(".//w:t", ns):
                                t.text = ""

                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")

                zout.writestr(item, data)

    os.remove(temp_docx)